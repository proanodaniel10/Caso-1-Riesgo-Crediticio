import json, math
from pathlib import Path
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter, FuncFormatter
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as I, Pt as P
from pptx.dml.color import RGBColor as PC
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parents[1]
J=json.loads((ROOT/'resultados/resultados_dashboard.json').read_text(encoding='utf-8'))
F=J['final'];ED=J['eda_train'];D=J['data'];CV=J['comparacion_cv']
DOC=ROOT/'reportes';AS=ROOT/'assets';DOC.mkdir(exist_ok=True);AS.mkdir(exist_ok=True)
CS=['#227c77','#d89a4e','#5b8cc8','#9171b1']
MC={'Regresión Logística':'#5b8cc8','Árbol de Decisión':'#d89a4e','Random Forest':'#9171b1','Gradient Boosting':'#227c77'}

def percent(n,dec=2):return f'{100*n:.{dec}f}%'
def number(n):return f'{n:,}'.replace(',','.')
def money(n):return 'US$ '+number(n)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.titlesize':12,'axes.titleweight':'bold','axes.spines.top':False,'axes.spines.right':False,'savefig.dpi':180})

def save(name):
    plt.tight_layout();plt.savefig(AS/name,facecolor='white',bbox_inches='tight');plt.close()

# Figures
fig,ax=plt.subplots(figsize=(7.9,2.7))
vals=[D['target']['0'],D['target']['1']]
ax.barh(['No incumple (0)','Incumple (1)'],vals,color=['#227c77','#d89a4e'],height=.52)
ax.invert_yaxis();ax.set_xlim(0,max(vals)*1.20)
for i,v in enumerate(vals):ax.text(v+300,i,number(v)+' · '+percent(v/sum(vals)),va='center',fontsize=10,weight='bold')
ax.set_xlabel('Solicitudes');ax.set_title('Distribución de la variable objetivo — dataset completo',loc='left',pad=13)
save('01_distribucion.png')

fig,ax=plt.subplots(figsize=(8.1,3.4))
names=[x['modelo'] for x in CV];vals=[x['cv_media']['pr_auc'] for x in CV]
y=np.arange(len(names));ax.barh(y,vals,color=[MC[n] for n in names],height=.52)
ax.set_yticks(y,names);ax.invert_yaxis();ax.xaxis.set_major_formatter(PercentFormatter(1));ax.set_xlim(0,max(vals)*1.27)
for i,v in enumerate(vals):ax.text(v+.002,i,percent(v),va='center',fontsize=9,weight='bold')
ax.axvline(ED['target_rate'],c='#b9bfc5',ls='--',lw=1.5,label='Prevalencia clase positiva (referencia)');ax.legend(loc='lower right',fontsize=8)
ax.set_title('PR-AUC / Average Precision — media CV (5 folds)',loc='left',pad=12)
save('02_comparacion_ap.png')

fig,axes=plt.subplots(1,2,figsize=(10,4.0))
for name in J['metodologia']['modelos']:
    r=J['modelos'][name]
    roc=np.array(r['curvas_oof']['roc']);pr=np.array(r['curvas_oof']['pr'])
    axes[0].plot(roc[:,0],roc[:,1],lw=2.0,label=name,c=MC[name])
    axes[1].plot(pr[:,0],pr[:,1],lw=2.0,label=name,c=MC[name])
axes[0].plot([0,1],[0,1],ls='--',c='#b8c4ce',lw=1)
axes[1].axhline(ED['target_rate'],ls='--',c='#b8c4ce',lw=1)
axes[0].set(title='ROC — OOF TRAIN',xlabel='FPR',ylabel='TPR',xlim=(0,1),ylim=(0,1))
axes[1].set(title='Precision-Recall — OOF TRAIN',xlabel='Recall',ylabel='Precision',xlim=(0,1),ylim=(0,1))
fig.legend(*axes[0].get_legend_handles_labels(),loc='lower center',ncol=2,bbox_to_anchor=(.5,-.04),fontsize=8)
fig.subplots_adjust(bottom=.25,wspace=.3)
save('03_curvas_oof.png')

cm=F['confusion_test'];fig,ax=plt.subplots(figsize=(5.4,4.0))
m=np.array([[cm['tn'],cm['fp']],[cm['fn'],cm['tp']]])
ax.imshow([[0,.6],[.6,0]],cmap='Oranges',vmin=0,vmax=1)
for i in range(2):
 for j in range(2):
    ax.text(j,i,number(m[i,j]),ha='center',va='center',fontsize=24,weight='bold',color='#213f50')
ax.set_xticks([0,1],['Predice 0','Predice 1']);ax.set_yticks([0,1],['Real 0','Real 1'])
ax.set_title('Matriz de confusión · TEST final (umbral 0,50)',pad=14)
ax.tick_params(length=0);save('04_matriz_test.png')

fig,ax=plt.subplots(figsize=(7.5,3.5))
for key,color,label in [('recall','#d89a4e','Recall'),('precision','#227c77','Precision'),('f1','#5b8cc8','F1')]:
    th=J['modelos'][F['modelo']]['umbrales_oof']
    ax.plot([x['umbral'] for x in th],[x[key] for x in th],label=label,marker='.',c=color,lw=2.0)
ax.axvline(.5,c='#adbcc6',ls='--',label='Umbral del test = 0,50');ax.set_ylim(0,1)
ax.set_xlabel('Umbral de clasificación');ax.set_ylabel('Valor de métrica');ax.set_title('Impacto de umbral en TRAIN (predicciones OOF)',loc='left');ax.legend(fontsize=8,loc='upper right')
save('05_umbral_oof.png')

fig,ax=plt.subplots(figsize=(7.9,3.4))
keys=['roc_auc','pr_auc','recall','precision','f1']
xs=np.arange(len(keys));w=.36
ax.bar(xs-w/2,[F['cv_media'][k] for k in keys],width=w,color='#227c77',label='Media de CV')
ax.bar(xs+w/2,[F['metricas_test'][k] for k in keys],width=w,color='#d89a4e',label='TEST')
ax.set_xticks(xs,['ROC-AUC','PR-AUC','Recall','Precision','F1']);ax.set_ylim(0,1);ax.yaxis.set_major_formatter(PercentFormatter(1))
ax.set_title('Modelo seleccionado: CV frente a TEST',loc='left');ax.legend()
save('06_cv_test.png')

# Simulación académica en USD, SOLO TRAIN OOF. Montos ficticios, no pérdidas observadas.
W_FN,W_FP=500,100
fig,ax=plt.subplots(figsize=(7.6,3.8))
names=J['metodologia']['modelos']; xs=np.arange(len(names)); barw=.36
for shift,cut,label,hatch in [(-barw/2,.15,'Umbral 0,15',''),(barw/2,.50,'Umbral 0,50','//')]:
    fn=[];fp=[]
    for name in names:
        row=next(r for r in J['modelos'][name]['umbrales_oof'] if abs(r['umbral']-cut)<1e-8)
        fn.append(row['confusion']['fn']*W_FN);fp.append(row['confusion']['fp']*W_FP)
    ax.bar(xs+shift,fn,width=barw,color='#cf8742',alpha=.90,label='FN × US$ 500 · '+label)
    ax.bar(xs+shift,fp,bottom=fn,width=barw,color='#4986aa',alpha=.9,label='FP × US$ 100 · '+label)
ax.set_xticks(xs,['Logística','Árbol','Random Forest','Gradient Boosting'])
ax.set_ylabel('Costo hipotético (millones de USD)')
ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _pos: f'{y/1_000_000:.1f}'.replace('.', ',')))
ax.set_title('Pérdida ilustrativa según umbral · 40.000 predicciones OOF TRAIN',loc='left')
ax.tick_params(axis='x',labelsize=8)
from matplotlib.patches import Patch
ax.legend(handles=[Patch(color='#cf8742',label='Falsos negativos (FN × US$ 500)'),Patch(color='#4986aa',label='Falsos positivos (FP × US$ 100)')],fontsize=8)
ax.text(.01,-.22,'Dos barras por algoritmo: izquierda = umbral 0,15 / derecha = umbral 0,50. Importes hipotéticos en USD, no pérdidas bancarias observadas.',transform=ax.transAxes,fontsize=8,color='#596d79')
save('07_escenarios_costo_oof.png')


# ===== DOCX =====
doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.73);sec.bottom_margin=Inches(.67);sec.left_margin=Inches(.84);sec.right_margin=Inches(.84)
styles=doc.styles
styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(9)
styles['Normal'].paragraph_format.space_after=Pt(5)
for h,size in [('Title',25),('Heading 1',16),('Heading 2',12),('Heading 3',10)]:
 s=styles[h];s.font.name='Aptos Display';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(22,68,78)
 s.paragraph_format.space_before=Pt(9);s.paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=footer.add_run('UCE · Caso 48  |  Análisis académico reproducible  |  Octubre de 2026');rr.font.size=Pt(8);rr.font.color.rgb=RGBColor(115,131,141)
header=sec.header.paragraphs[0];header.text='INGENIERÍA EN SISTEMAS DE INFORMACIÓN  /  MINERÍA DE DATOS'
for r in header.runs:r.font.size=Pt(7);r.font.color.rgb=RGBColor(125,146,155)

def p(txt='',style=None):return doc.add_paragraph(txt,style=style)
def hd(txt,n=1):return doc.add_heading(txt,level=n)
def add_table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,h in enumerate(headers):t.rows[0].cells[j].text=str(h)
 for row in rows:
  cells=t.add_row().cells
  for j,x in enumerate(row):cells[j].text=str(x)
 for row in t.rows:
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2)
    for r in pp.runs:r.font.size=Pt(8)
 return t

def bullet(text):
 doc.add_paragraph(text,style='List Bullet')
def fig_doc(file,width=5.7,caption=None):
 pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pp.add_run().add_picture(str(AS/file),width=Inches(width))
 if caption:
  cp=p(caption);cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
  for r in cp.runs:r.font.italic=True;r.font.size=Pt(8);r.font.color.rgb=RGBColor(99,122,135)
def page():doc.add_page_break()

p('INFORME TÉCNICO DE EXAMEN COMPLEXIVO')
doc.add_heading('Caso 48',0)
p('Predicción de incumplimiento crediticio mediante clasificación supervisada')
p('Universidad Central del Ecuador  ·  Facultad de Ingeniería y Ciencias Aplicadas')
p('Carrera: Ingeniería en Sistemas de Información')
p('Componente: Minería de Datos  |  Subcomponente: Clasificación supervisada y evaluación de modelos')
p('Documento basado en el caso oficial y en la base de datos sintética recibida. Resultados obtenidos ejecutando el script reproducible del proyecto. No representa un modelo de producción.')
add_table(['Indicador','Resultado'],[
 ('Solicitudes analizadas',number(D['rows_used'])),('Variables originales',D['columns_original']),
 ('Predictores utilizados',len(D['predictors'])),('Incumplimientos (clase 1)',f"{number(D['target']['1'])} ({percent(D['target_rate'])})"),
 ('Modelo elegido únicamente por CV',F['modelo']),('PR-AUC media CV',percent(F['cv_media']['pr_auc'])),
 ('PR-AUC TEST',percent(F['metricas_test']['pr_auc'])),('Recall TEST con umbral 0,50',percent(F['metricas_test']['recall']))])
p('Resumen ejecutivo',style='Heading 2')
p('Se entrenaron cuatro clasificadores con preprocesamiento dentro de Pipeline/ColumnTransformer y StratifiedKFold de cinco particiones. El mejor resultado de PR-AUC media de validación correspondió a Gradient Boosting. Aunque muestra capacidad de ordenación superior al azar (ROC-AUC TEST 0,6756), al aplicar el umbral fijo 0,50 no identificó ningún incumplimiento en TEST. Esta limitación es central y obliga a distinguir capacidad de ranking de eficacia operativa en un umbral dado.')

page();hd('1. Objetivo, alcance y alineación al caso')
p('El documento oficial solicita predecir default_12m (1 = incumplimiento en los siguientes 12 meses; 0 = no incumplimiento), comparar cuatro clasificadores y discutir falsos positivos/negativos sin utilizar TEST para elegir el modelo. El dataset es sintético y solo tiene fines didácticos (PDF: p. 1 y p. 2).')
p('Se cumplen las doce instrucciones: inspección inicial, EDA, definición de X/y, holdout 80/20, pipelines, cuatro modelos, validación cruzada estratificada de cinco folds, métricas, selección por CV, test independiente, comparación CV/test y conclusiones responsables.')
hd('2. Auditoría de datos',1)
p(f"Se recibieron {number(D['rows_original'])} registros y {D['columns_original']} columnas: 14 variables predictoras, un identificador y una variable objetivo. No se detectaron valores faltantes, duplicados completos ni identificadores repetidos. Se excluyó applicant_id de X por tratarse de una clave arbitraria, y default_12m por ser la variable objetivo.")
add_table(['Control','Hallazgo'],[
 ('Dimensiones',f"{number(D['rows_used'])} filas × {D['columns_original']} columnas"),
 ('Valores faltantes',number(D['missing_total'])),('Duplicados completos',number(D['duplicates_full'])),
 ('Identificadores repetidos',number(D['duplicates_id'])),
 ('No incumple (0)',number(D['target']['0'])),('Incumple (1)',number(D['target']['1']))])
fig_doc('01_distribucion.png',5.6,'Figura 1. Desbalance del objetivo en los datos recibidos (inspección inicial).')
p('Una accuracy del 94,52 % sería alcanzable clasificando siempre como no incumplidor, con Recall de clase 1 igual a cero. Por eso la accuracy no puede ser el criterio exclusivo del ejercicio.')

page();hd('3. Exploración de TRAIN')
p('Se separó el conjunto de prueba antes de realizar exploración supervisada. La EDA usa 40.000 ejemplos de entrenamiento; 2.194 son positivos. Se revisaron estadísticas descriptivas, distribuciones numéricas, frecuencias categóricas, relaciones por clase y outliers por IQR. Los valores extremos se describieron, sin eliminarlos automáticamente.')
add_table(['Variable (TRAIN)','Media','Mediana','Atípicos IQR'],[
 (k, f"{ED['numeric'][k]['mean']:.2f}", f"{ED['numeric'][k]['median']:.2f}",number(ED['numeric'][k]['outliers_iqr']))
 for k in ['age','monthly_income_usd','debt_to_income_ratio','credit_utilization','late_payments_12m','savings_usd']])
p('Como ejemplos descriptivos no causales, el incumplimiento en TRAIN es del '+percent(next(x['tasa'] for x in ED['categories']['employment_type'] if x['categoria']=='independiente'))+' entre trabajadores independientes y del '+percent(next(x['tasa'] for x in ED['categories']['employment_type'] if x['categoria']=='dependiente'))+' entre trabajadores dependientes. Estos porcentajes no justifican decisiones individuales ni implican causalidad.')
hd('4. Diseño metodológico')
add_table(['Etapa','Protocolo'],[
 ('Holdout','80 % TRAIN (40.000) / 20 % TEST (10.000), estratificado y seed 42.'),
 ('Variables','14 predictores; applicant_id y default_12m excluidos de X.'),
 ('Transformaciones','Pipeline / ColumnTransformer: imputación mediana numérica, imputación moda + One-Hot categórica.'),
 ('Escalado','StandardScaler solo para Regresión Logística.'),
 ('CV','StratifiedKFold(5, shuffle=True, random_state=42), exclusivamente TRAIN.'),
 ('Selección','Mayor PR-AUC = Average Precision media CV; desempate por Recall CV.'),
 ('Umbral final','0,50, preestablecido; variaciones solo explicativas con predicciones OOF TRAIN.'),
 ('Evaluación','Única evaluación final en TEST del modelo seleccionado por CV.')])

page();hd('5. Algoritmos y validación cruzada')
add_table(['Modelo','Parámetros elegidos previamente'],[
 ('Regresión Logística','class_weight=balanced, max_iter=1300 y StandardScaler.'),
 ('Árbol de Decisión','max_depth=6, min_samples_leaf=25, class_weight=balanced.'),
 ('Random Forest','110 árboles, min_samples_leaf=5, class_weight=balanced_subsample.'),
 ('Gradient Boosting','HistGradientBoostingClassifier, 110 iteraciones, max_leaf_nodes=15, learning_rate=0,07.')])
p('HistGradientBoostingClassifier es una implementación del algoritmo Gradient Boosting optimizada mediante histogramas. Los hiperparámetros son fijos y no se eligieron mirando TEST.')
add_table(['Modelo','Precision CV','Recall CV','F1 CV','ROC-AUC CV','PR-AUC CV'],[
 (r['modelo'],percent(r['cv_media']['precision'],1),percent(r['cv_media']['recall'],1),
  percent(r['cv_media']['f1'],1),percent(r['cv_media']['roc_auc'],1),percent(r['cv_media']['pr_auc'],1)) for r in CV])
fig_doc('02_comparacion_ap.png',5.8,'Figura 2. PR-AUC (AP) media entre cinco folds, calculada solo sobre TRAIN.')
p('Resultado: Gradient Boosting fue seleccionado por una PR-AUC media de '+percent(F['cv_media']['pr_auc'])+' (desviación estándar '+percent(F['cv_desv']['pr_auc'])+'). La prevalencia del objetivo, cerca de 5,5 %, sirve de referencia interpretativa para AP; no debe confundirse con el desempeño de un modelo real.')

page();hd('6. Curvas de clasificación sobre TRAIN')
p('Para generar curvas comparables de los cuatro modelos se utilizaron predicciones out-of-fold (OOF). Cada fila TRAIN fue puntuada por un modelo que no la utilizó durante el ajuste. Este análisis no participa en la selección posterior del modelo en TEST.')
fig_doc('03_curvas_oof.png',6.35,'Figura 3. Curvas ROC y Precision–Recall OOF de los cuatro algoritmos.')
hd('7. Evaluación final en TEST')
p(f"Se reentrenó el pipeline de {F['modelo']} con los 40.000 registros de TRAIN y se evaluó una sola vez sobre los 10.000 ejemplos reservados. La clase positiva tuvo 548 incumplimientos reales.")
add_table(['Métrica','TEST final'],[(k,percent(F['metricas_test'][k])) for k in ['precision','recall','f1','roc_auc','pr_auc','accuracy']])
fig_doc('04_matriz_test.png',4.7,'Figura 4. Matriz de confusión del modelo seleccionado, TEST, umbral fijo 0,50.')
p(f"VN={number(cm['tn'])}, FP={number(cm['fp'])}, FN={number(cm['fn'])}, VP={number(cm['tp'])}. El modelo predijo todos los casos como clase 0, por lo que Precision, Recall y F1 son 0 %. Este comportamiento no debe presentarse como un sistema de detección eficaz.")

page();hd('8. Umbral de decisión y costos de errores')
p('La comparación de falsos positivos (marcar como incumplidor a quien pagaría) y falsos negativos (no identificar a quien incumpliría) debe reconocer que el enunciado NO facilita costos monetarios reales. Un falso negativo puede implicar exposición al impago; un falso positivo podría producir revisión adicional o perjuicios si se automatizara una decisión. Por ello, se simulan importes en dólares estadounidenses (USD), explícitamente hipotéticos, sin atribuir al ejercicio pérdidas financieras verificadas.')
fig_doc('05_umbral_oof.png',5.8,'Figura 5. Sensibilidad pedagógica a distintos umbrales calculada exclusivamente en TRAIN OOF.')
th=J['modelos'][F['modelo']]['umbrales_oof']
add_table(['Umbral TRAIN OOF','Precision','Recall','FP','FN'],[
 (f"{x['umbral']:.2f}",percent(x['precision']),percent(x['recall']),number(x['confusion']['fp']),number(x['confusion']['fn']))
 for x in th if round(x['umbral'],2) in [.05,.10,.15,.20,.50]])
p('Al reducir el umbral pueden detectarse más incumplimientos, pero suelen aumentar los falsos positivos. Los datos de esta tabla provienen solo de TRAIN con predicciones fuera de fold; NO son una segunda evaluación sobre TEST ni resultados de un sistema en producción.')
page();hd('8.1 Análisis cuantitativo de costos hipotéticos en USD',n=2)
p('Supuestos monetarios SOLO para demostración: C_FN = US$ 500 por cada incumplimiento no detectado; C_FP = US$ 100 por cada buen pagador revisado o rechazado innecesariamente. Son cantidades expresadas en USD, pero NO provienen del CSV ni del enunciado y NO son pérdidas reales de una institución financiera.')
p('COSTO TOTAL HIPOTÉTICO (USD) = FN × C_FN(USD) + FP × C_FP(USD). Se utilizan los conteos de 40.000 predicciones fuera de fold (OOF) de TRAIN; el conjunto TEST permanece aislado y el modelo final continúa seleccionado por la mayor PR-AUC media en CV.')
add_table(['Modelo','Umbral','FN','FP','Costo FN (USD)','Costo FP (USD)','Total (USD)'],[
 (name,f'{cut:.2f}',number(row['confusion']['fn']),number(row['confusion']['fp']),
  money(500*row['confusion']['fn']),money(100*row['confusion']['fp']),
  money(500*row['confusion']['fn']+100*row['confusion']['fp']))
 for name in J['metodologia']['modelos'] for cut in [.15,.5]
 for row in J['modelos'][name]['umbrales_oof'] if abs(row['umbral']-cut)<1e-8
])
fig_doc('07_escenarios_costo_oof.png',5.65,'Figura 6. Costo simulado en USD, solo TRAIN OOF. FN: US$ 500 por error; FP: US$ 100 por error. No representa pérdidas reales.')
p('En Gradient Boosting el umbral 0,15 genera FN=2.003 y FP=812 (costo hipotético US$ 1.082.700); el umbral 0,50 genera FN=2.194 y FP=0 (costo hipotético US$ 1.097.000). Esta comparación explica la compensación entre errores, NO valida económicamente ni recomienda alterar el umbral del TEST existente. Para decisiones reales se necesitarían costos observados, políticas de revisión y otra evaluación independiente.')
page()

hd('9. Generalización y posible sobreajuste')
p('PR-AUC TRAIN interno promedio = '+percent(F['train_folds_media']['pr_auc'])+'; PR-AUC CV media = '+percent(F['cv_media']['pr_auc'])+'; PR-AUC TEST = '+percent(F['metricas_test']['pr_auc'])+'. La brecha TRAIN interno frente a CV de '+percent(F['gap_ap_train_cv'])+' puntos porcentuales sugiere sobreajuste o una discrepancia importante entre ajuste y validación. TEST no evidencia una caída adicional grande respecto de CV, pero ello no elimina el problema del umbral 0,50.')

page();hd('10. Conclusiones y recomendaciones')
for x in [
 'Se verificaron 50.000 observaciones y 16 columnas completas, sin nulos ni duplicados; el objetivo muestra un desbalance pronunciado (5,48 % positivos).',
 'Los cuatro algoritmos se compararon con cinco folds estratificados, preprocesamiento dentro de los pipelines y TEST aislado de la selección.',
 'Gradient Boosting encabezó la PR-AUC media de CV (11,38 %), pero su Recall y F1 a umbral 0,50 fueron 0 % en TRAIN y TEST; la accuracy no demuestra detección efectiva.',
 'La diferencia entre PR-AUC de TRAIN interno y CV amerita cautela por posible sobreajuste. La utilidad del clasificador depende de umbral, costos de errores y condiciones de aplicación.',
 'Se simularon costos hipotéticos en USD (FN = US$ 500 y FP = US$ 100 por error), exclusivamente en TRAIN OOF. Antes de cualquier uso real habría que obtener costos observados, definir y validar un umbral en TRAIN y realizar una evaluación independiente nueva.',
 'El dataset es sintético, carece de evidencia de generalización temporal y no autoriza decisiones reales de aprobación o rechazo automático de créditos.'
]:bullet(x)
hd('11. Limitaciones y responsabilidad académica')
p('Los resultados dependen de una base artificial, de una única división estratificada y de hiperparámetros predefinidos. Ninguna métrica revela causalidad. No se evaluaron cambios económicos futuros, deriva de datos, sesgos de representación ni equidad por grupos protegidos. Una implementación real requeriría auditorías independientes, trazabilidad, supervisión humana y requisitos regulatorios aplicables.')
hd('12. Referencias y fuentes de evidencia')
bullet('Universidad Central del Ecuador. Caso 48: Predicción de incumplimiento crediticio, documento oficial proporcionado (2 páginas), instrucciones 1–12 y diccionario de variables.')
bullet('Archivo riesgo_crediticio_bigdata.csv proporcionado para el ejercicio; 50.000 filas, 16 columnas; dataset sintético didáctico.')
bullet('Código reproducible: src/entrenar.py; salidas tabulares en resultados/ y dashboard agregado en dashboard/datos.js.')
bullet('Documentación de scikit-learn para Pipeline, ColumnTransformer, StratifiedKFold, HistGradientBoostingClassifier y métricas.')
bullet('GitHub Docs: GitHub Pages permite HTML/CSS/JS estático, no ejecuta Python del lado del servidor.')

page();hd('Anexo A. Correspondencia con las 12 instrucciones')
mapa=[
(1,'Control de dimensiones, tipos, duplicados y nulos','Auditoría del dataset / src/entrenar.py'),
(2,'EDA de clase objetivo, descriptivos, frecuencias, outliers y relaciones','Dashboard Exploración; tablas train'),
(3,'Separar X/y y excluir applicant_id','Preparación X = 14 predictores, y = default_12m'),
(4,'Holdout 80/20 estratificado y TEST reservado','Método / src/entrenar.py'),
(5,'ColumnTransformer/Pipeline, One-Hot, escalado','processor() / build_models()'),
(6,'Cuatro algoritmos de clasificación','Modelos de src/entrenar.py'),
(7,'StratifiedKFold de 5 folds solo TRAIN','Ciclo de CV y métricas_por_fold.csv'),
(8,'Precision, Recall, F1, ROC-AUC, PR-AUC, matriz de confusión','Dashboard CV/TEST y gráficos'),
(9,'Selección mediante CV y análisis FP/FN','Ranking por AP; dashboard Costos y errores · P9; costos monetarios SIMULADOS en USD usando TRAIN OOF'),
(10,'Reentrenar final y evaluar una sola vez en TEST','modelo final y metricas_test_FINAL.csv'),
(11,'Comparar CV vs TEST, estabilidad y sobreajuste','cv_vs_test.csv y sección de análisis'),
(12,'Conclusiones, límites y condiciones para uso real','Informe: secciones 10 y 11')]
add_table(['N.º','Requerimiento del caso','Evidencia en el proyecto'],mapa)
p('Glosario para la defensa',style='Heading 2')
for a,b in [('OOF','Predicción fuera de fold; cada ejemplo se predice por un modelo que no lo entrenó.'),
 ('PR-AUC / AP','Aquí es Average Precision; no es idéntica al área trapezoidal bajo la curva PR.'),
 ('Precision','Porcentaje de positivos predichos que realmente son positivos.'),
 ('Recall','Porcentaje de incumplidores reales correctamente detectados.'),
 ('Falsos negativos','Personas que incumplen clasificadas como no incumplidoras.'),
 ('Data leakage','Uso indebido de información que no debería estar disponible durante entrenamiento/selección.')]:
 p(f'{a}: {b}')

report=DOC/'Informe_Tecnico_Caso_48_UCE.docx'
doc.save(report)
print('INFORME:',report)

# ===== Presentación PPTX para defensa =====
prs=Presentation();prs.slide_width=I(13.33);prs.slide_height=I(7.5)
NBG=PC(16,33,51);INK=PC(26,51,66);TEAL=PC(27,135,124);WHITE=PC(255,255,255);MUTED=PC(105,124,134);GOLD=PC(217,154,77)

def rect(slide,x,y,w,h,fill,line=None,radius=False):
 shp=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,I(x),I(y),I(w),I(h))
 shp.fill.solid();shp.fill.fore_color.rgb=fill
 shp.line.fill.background() if line is None else None
 return shp

def txt(slide,text,x,y,w,h,size=22,bold=False,color=INK,align=None):
 sh=slide.shapes.add_textbox(I(x),I(y),I(w),I(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=I(.02);tf.margin_right=I(.02);tf.margin_top=I(.01);tf.margin_bottom=I(.01)
 pp=tf.paragraphs[0];pp.text=text;pp.font.name='Aptos';pp.font.size=P(size);pp.font.bold=bold;pp.font.color.rgb=color
 if align is not None:pp.alignment=align
 return sh

def newslide(k,title,subtitle=None,dark=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,13.33,7.5,NBG if dark else WHITE)
 rect(s,.45,.50,.065,.35,TEAL)
 txt(s,'CASO / 48    ·    MINERÍA DE DATOS',.67,.48,10,.25,10,True,PC(126,222,210) if dark else TEAL)
 txt(s,title,.67,.98,11.9,.72,29,True,WHITE if dark else INK)
 if subtitle:txt(s,subtitle,.67,1.76,11.9,.55,14,False,PC(188,209,222) if dark else MUTED)
 rect(s,.65,7.12,12,.012,PC(62,88,105) if dark else PC(223,230,234))
 txt(s,f'UNIVERSIDAD CENTRAL DEL ECUADOR   ·   EVALUACIÓN REPRODUCIBLE',.66,7.19,10,.16,8,False,PC(183,202,214) if dark else MUTED)
 txt(s,f'{k:02d} / 13',11.95,7.16,.7,.19,9,True,PC(183,202,214) if dark else MUTED)
 return s

def big(s,n,label,x,y,w=3.65):
 rect(s,x,y,w,1.25,PC(239,248,246),radius=True);txt(s,n,x+.23,y+.15,w-.46,.55,27,True,TEAL)
 txt(s,label,x+.23,y+.83,w-.46,.2,11,False,MUTED)

def point(s,head,body,x,y,w=11.5):
 rect(s,x,y,.08,.48,TEAL);txt(s,head,x+.23,y-.05,w,.3,17,True,INK);txt(s,body,x+.23,y+.32,w,.45,13,False,MUTED)

def chart(s,name,x,y,w,h):s.shapes.add_picture(str(AS/name),I(x),I(y),width=I(w),height=I(h))

s=newslide(1,'Predicción de incumplimiento crediticio','Examen complexivo UCE · Cuatro modelos · Validación cruzada · Dashboard',True)
txt(s,'DE DATOS A EVIDENCIA',.75,2.85,10,.5,17,True,PC(121,227,208))
txt(s,'50.000 solicitudes\n14 variables predictoras\nEvaluación independiente en TEST',.75,3.42,10,2.4,27,True,WHITE)

s=newslide(2,'¿Cuál es el problema?','Predecir default_12m: 1 incumple en 12 meses; 0 no incumple.')
point(s,'Aprendizaje supervisado','Existe una variable objetivo etiquetada para cada solicitud.',.82,2.45)
point(s,'Clasificación binaria','La predicción es una clase, no un importe monetario.',.82,3.65)
point(s,'Restricción del examen','Seleccionar el modelo sin mirar TEST; reportar efectos de FP y FN.',.82,4.85)

s=newslide(3,'Inspección del dataset','CSV oficial del ejercicio · Fuente sintética y didáctica')
big(s,'50.000','Registros',.82,2.30);big(s,'16','Columnas',4.85,2.30);big(s,'0','Nulos / duplicados',8.88,2.30)
chart(s,'01_distribucion.png',1.0,4.03,10.8,2.45)

s=newslide(4,'El reto es el desbalance','Solo el 5,48 % de los solicitantes tiene incumplimiento.')
chart(s,'01_distribucion.png',.87,2.20,7.0,3.1)
rect(s,8.5,2.28,3.88,3.15,PC(255,246,233),radius=True)
txt(s,'94,52 %',8.87,2.72,3.2,.7,38,True,PC(150,89,29))
txt(s,'Accuracy del modelo trivial\nque siempre predice «no incumple».\nRecall = 0 %.',8.90,3.64,3.14,1.2,17,False,INK)

s=newslide(5,'Metodología sin fuga de información','Una secuencia reproducible con seed 42.')
for i,(a,b) in enumerate([('01','X/y, excluir ID'),('02','80 % TRAIN / 20 % TEST'),('03','Pipeline + One-Hot'),('04','StratifiedKFold(5)'),('05','Selección por AP CV'),('06','Evaluación única TEST')]):
 col=i%3; row=i//3;x=.85+4.15*col;y=2.25+1.84*row
 rect(s,x,y,3.78,1.45,PC(238,248,246),radius=True)
 txt(s,a,x+.18,y+.12,.66,.38,21,True,TEAL);txt(s,b,x+.18,y+.65,3.44,.52,15,True,INK)

s=newslide(6,'Comparación de los cuatro algoritmos','Resultado de la validación cruzada: PR-AUC como criterio principal.')
chart(s,'02_comparacion_ap.png',.9,2.12,8.2,4.12)
rect(s,9.2,2.45,3.1,2.52,PC(235,247,243),radius=True)
txt(s,'GANADOR CV',9.57,2.75,2.48,.36,12,True,TEAL)
txt(s,'Gradient\nBoosting',9.55,3.29,2.59,.95,22,True,INK)

s=newslide(7,'Curvas de los cuatro modelos','ROC y Precision-Recall fuera de fold: nunca se usa TEST para elegir.')
chart(s,'03_curvas_oof.png',.75,2.1,11.88,4.62)

s=newslide(8,'La gran limitación del umbral 0,50','Gradient Boosting gana por PR-AUC, pero su Recall a ese umbral es 0 %.')
chart(s,'05_umbral_oof.png',.75,2.22,8.15,4.32)
rect(s,9.12,2.49,3.25,3.2,PC(255,244,235),radius=True)
txt(s,'TRAIN OOF',9.40,2.81,2.55,.34,12,True,PC(148,91,40))
txt(s,'Cambiar el umbral altera el balance entre FP y FN.',9.40,3.41,2.55,1.65,18,True,INK)

s=newslide(9,'Punto 9: costo de equivocarse','Supuestos en USD, NO pérdidas reales: FN = US$ 500 y FP = US$ 100 · TRAIN OOF')
chart(s,'07_escenarios_costo_oof.png',.85,2.18,8.6,4.10)
rect(s,9.55,2.42,2.85,3.65,PC(238,248,246),radius=True)
txt(s,'GRADIENT BOOSTING',9.85,2.75,2.40,.44,12,True,TEAL)
txt(s,'Umbral 0,15: US$ 1.082.700\nUmbral 0,50: US$ 1.097.000',9.85,3.30,2.25,1.40,17,True,INK)
txt(s,'Costos ilustrativos, seleccion final por PR-AUC CV. TEST aislado.',9.85,5.0,2.25,.90,12,False,MUTED)

s=newslide(10,'TEST: evaluación independiente','10.000 solicitudes reservadas · 548 incumplimientos reales.')
chart(s,'04_matriz_test.png',.8,2.19,6.55,4.42)
for i,(v,lab) in enumerate([('0 %','Recall final'),(percent(F['metricas_test']['pr_auc']),'PR-AUC final'),(percent(F['metricas_test']['roc_auc']),'ROC-AUC final')]):
 big(s,v,lab,8.10,2.24+i*1.38,3.55)

s=newslide(11,'¿Existe sobreajuste?','La PR-AUC interna de entrenamiento supera ampliamente la de validación.')
chart(s,'06_cv_test.png',.9,2.21,8.8,3.95)
txt(s,'PR-AUC TRAIN interno: '+percent(F['train_folds_media']['pr_auc'])+'\nPR-AUC media CV: '+percent(F['cv_media']['pr_auc'])+'\nPR-AUC TEST: '+percent(F['metricas_test']['pr_auc']),9.6,2.7,2.8,2.6,17,True,INK)

s=newslide(12,'Conclusión y uso responsable','El proyecto cumple el protocolo; los límites importan tanto como las métricas.')
point(s,'Hallazgo','La capacidad de ranking es modesta; el umbral 0,50 no detectó positivos.',.82,2.38)
point(s,'Mejoras necesarias','Umbral definido con TRAIN/CV, calibración, costos de errores y validación externa.',.82,3.75)
point(s,'Límite de aplicación','Dataset sintético; no usar como aprobador o rechazador automático de crédito.',.82,5.12)

s=newslide(13,'Preguntas que debes poder responder','Guía breve para exponer sin depender de leer las diapositivas.')
for i,(a,b) in enumerate([
 ('¿Por qué PR-AUC?','Por el desbalance de 5,48 % positivos; accuracy engaña.'),
 ('¿Qué es CV de 5 folds?','Cinco entrenamientos con particiones estratificadas de TRAIN.'),
 ('¿Por qué no usaste TEST en la selección?','Evita optimismo y fuga de información al evaluar.'),
 ('¿Por qué Gradient Boosting si Recall es cero?','Fue el primero por AP CV; ranking ≠ detección al umbral 0,50.'),
 ('¿Cuál es el siguiente paso?','Definir y validar umbral solo en TRAIN; probar en otro test independiente.')]):
  y=2.22+i*.86;txt(s,a,.95,y,5.9,.31,15,True,TEAL);txt(s,b,6.82,y,5.7,.6,12,False,INK)
pp=DOC/'Defensa_Caso_48_UCE.pptx';prs.save(pp)
print('DEFENSA:',pp)
print('FIGURAS:',len(list(AS.glob('*.png'))))
